"""Извлечение текста из документов (DOCX/PDF/XLSX). OCR не в MVP (Q6)."""
import subprocess
from pathlib import Path
import docx
from openpyxl import load_workbook


def extract_docx(path: str) -> str:
    """Текст из параграфов И таблиц (в договорах CargoTech тело часто в таблице)."""
    d = docx.Document(path)
    parts = []
    for p in d.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


def extract_pdf(path: str) -> str:
    """PDF с текстовым слоём через poppler pdftotext (-layout сохраняет структуру)."""
    out = subprocess.run(
        ["pdftotext", "-layout", "-enc", "UTF-8", path, "-"],
        capture_output=True, text=True,
    )
    return out.stdout


def extract_xlsx(path: str) -> str:
    """Excel (выгрузки 1С) -> табличный текст. В RAG НЕ векторизуем, только для справки."""
    wb = load_workbook(path, data_only=True)
    rows = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                rows.append(" | ".join(cells))
    return "\n".join(rows)


def extract(path: str) -> str:
    ext = Path(path).suffix.lower()
    if ext == ".docx":
        return extract_docx(path)
    if ext == ".pdf":
        return extract_pdf(path)
    if ext == ".xlsx":
        return extract_xlsx(path)
    raise ValueError(f"Формат {ext} не поддерживается в MVP")
